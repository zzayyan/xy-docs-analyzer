"""Every supported format converts into chunks, and nothing else does. Runs real Docling; no LLM or Neo4j."""
import docx
import openpyxl
import pptx
import pytest
from docling.exceptions import ConversionError
from PIL import Image, ImageDraw, ImageFont

from app.ingest.convert import convert


def test_a_format_outside_the_supported_list_is_not_converted(tmp_path):
    path = tmp_path / "notes.adoc"
    path.write_text("= Notes\n\nDenda 2% per bulan.\n", encoding="utf-8")
    with pytest.raises(ConversionError, match="not allowed"):
        convert(str(path), set())


def write_docx(path):
    doc = docx.Document()
    doc.add_heading("Pembayaran", level=1)
    doc.add_paragraph("Denda keterlambatan 2% per bulan.")
    doc.save(path)


def write_html(path):
    path.write_text("<html><body><h1>Pembayaran</h1><p>Denda keterlambatan 2% per bulan.</p></body></html>", encoding="utf-8")


def write_md(path):
    path.write_text("# Pembayaran\n\nDenda keterlambatan 2% per bulan.\n", encoding="utf-8")


@pytest.mark.parametrize("suffix, write", [(".docx", write_docx), (".html", write_html), (".md", write_md)])
def test_a_text_document_becomes_a_chunk_under_its_heading(tmp_path, suffix, write):
    path = tmp_path / f"kontrak{suffix}"
    write(path)
    _, _, chunks = convert(str(path), set())
    [chunk] = chunks
    assert chunk["headings"] == ["Pembayaran"]
    assert "Denda keterlambatan 2% per bulan." in chunk["text"]


def test_each_slide_becomes_a_chunk_under_its_title_with_the_slide_number_as_page(tmp_path):
    path = tmp_path / "rapat.pptx"
    deck = pptx.Presentation()
    for title, body in [("Rencana Q3", "Target pendapatan naik 20 persen."), ("Risiko", "Vendor terlambat mengirim server.")]:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    deck.save(path)

    _, _, chunks = convert(str(path), set())
    assert [(c["headings"], c["pages"]) for c in chunks] == [(["Rencana Q3"], [1]), (["Risiko"], [2])]
    assert "Vendor terlambat mengirim server." in chunks[1]["text"]


def write_xlsx(path):
    book = openpyxl.Workbook()
    sheet = book.active
    for row in [("paket", "harga"), ("Basic", 100000), ("Pro", 250000)]:
        sheet.append(row)
    book.save(path)


def write_csv(path):
    path.write_text("paket,harga\nBasic,100000\nPro,250000\n", encoding="utf-8")


@pytest.mark.parametrize("suffix, write", [(".xlsx", write_xlsx), (".csv", write_csv)])
def test_a_spreadsheet_row_keeps_its_values_together_in_a_chunk(tmp_path, suffix, write):
    path = tmp_path / f"harga{suffix}"
    write(path)
    _, _, chunks = convert(str(path), set())
    assert any("Pro" in c["text"] and "250000" in c["text"] for c in chunks)


@pytest.mark.parametrize("suffix", [".png", ".jpg"])
def test_text_in_a_picture_is_read_by_ocr(tmp_path, suffix):
    """Loads Docling's layout and OCR models, so it takes several seconds."""
    path = tmp_path / f"foto{suffix}"
    picture = Image.new("RGB", (1400, 700), "white")
    draw = ImageDraw.Draw(picture)
    draw.text((60, 60), "Pembayaran", fill="black", font=ImageFont.load_default(size=56))
    body = ImageFont.load_default(size=36)
    for i, line in enumerate(["Pembayaran dilakukan paling lambat tanggal 10 setiap bulan.",
                              "Keterlambatan dikenakan denda dua persen per bulan."]):
        draw.text((60, 200 + i * 60), line, fill="black", font=body)
    picture.save(path)

    _, _, chunks = convert(str(path), set())
    assert "denda dua persen" in " ".join(c["text"] for c in chunks).lower()
