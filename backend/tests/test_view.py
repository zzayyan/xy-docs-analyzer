"""GET /documents/{id}/view: the kept DoclingDocument as blocks a viewer marks by ref. Real Docling, store faked."""
import openpyxl
import pptx
import pytest
from fastapi.testclient import TestClient

import app.store as store_module
from app.api import app
from app.config import settings
from app.ingest.convert import convert, converted_path

client = TestClient(app)


@pytest.fixture
def ingest(monkeypatch, tmp_path):
    """Converts a file the way ingest does and keeps its document under the name given; returns its chunks."""
    monkeypatch.setattr(settings, "data_dir", tmp_path / "data")
    names = {}
    monkeypatch.setattr(store_module, "get_file", lambda doc_id: names.get(doc_id) and {"id": doc_id, "name": names[doc_id], "status": "done"})

    def run(path) -> list[dict]:
        doc, _, chunks = convert(str(path), set())
        converted_path("d1").parent.mkdir(parents=True, exist_ok=True)
        doc.save_as_json(converted_path("d1"))
        names["d1"] = path.name
        return chunks

    return run


def test_a_text_document_is_its_headings_and_paragraphs_with_the_cited_one_findable(ingest, tmp_path):
    path = tmp_path / "kontrak.html"
    path.write_text(
        "<html><body><h1>Pembayaran</h1><p>Denda keterlambatan 2% per bulan.</p>"
        "<script>alert('xss')</script></body></html>",
        encoding="utf-8",
    )
    [chunk] = ingest(path)

    view = client.get("/documents/d1/view").json()

    assert view["kind"] == "document"
    assert [(b["type"], b["text"]) for b in view["blocks"]] == [
        ("heading", "Pembayaran"),
        ("paragraph", "Denda keterlambatan 2% per bulan."),
    ], "the script is not part of the document"
    [cited] = [b for b in view["blocks"] if b["ref"] in chunk["refs"]]
    assert cited["text"] == "Denda keterlambatan 2% per bulan."


def test_slides_are_blocks_numbered_by_their_slide(ingest, tmp_path):
    path = tmp_path / "rapat.pptx"
    deck = pptx.Presentation()
    for title, body in [("Rencana Q3", "Target naik 20 persen."), ("Risiko", "Vendor terlambat.")]:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    deck.save(path)
    ingest(path)

    view = client.get("/documents/d1/view").json()

    assert view["kind"] == "slides"
    assert [(b["page"], b["type"], b["text"]) for b in view["blocks"]] == [
        (1, "heading", "Rencana Q3"),
        (1, "list_item", "Target naik 20 persen."),
        (2, "heading", "Risiko"),
        (2, "list_item", "Vendor terlambat."),
    ]


def test_each_sheet_is_its_name_and_table_and_a_cited_row_points_into_it(ingest, tmp_path):
    path = tmp_path / "harga.xlsx"
    book = openpyxl.Workbook()
    book.active.title = "Paket"
    for row in [("paket", "harga"), ("Basic", 100000), ("Pro", 250000)]:
        book.active.append(row)
    book.create_sheet("Diskon").append(("kode", "persen"))
    book["Diskon"].append(("HEMAT", 10))
    book.save(path)
    chunks = ingest(path)

    blocks = client.get("/documents/d1/view").json()["blocks"]

    assert [(b["type"], b["text"]) for b in blocks if b["type"] == "heading"] == [("heading", "Paket"), ("heading", "Diskon")]
    tables = {b["ref"]: b["rows"] for b in blocks if b["type"] == "table"}
    assert tables["#/tables/0"] == [["paket", "harga"], ["Basic", "100000"], ["Pro", "250000"]]
    refs = [ref for c in chunks for ref in c["refs"]]
    assert "#/tables/0/rows/2" in refs
    table, _, row = "#/tables/0/rows/2".rpartition("/rows/")
    assert tables[table][int(row)] == ["Pro", "250000"]


def test_view_is_404_for_an_unknown_document_or_one_converted_before_documents_were_kept(monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "data_dir", tmp_path)
    monkeypatch.setattr(store_module, "get_file", lambda doc_id: {"id": doc_id, "name": "old.docx", "status": "done"} if doc_id == "old" else None)

    assert client.get("/documents/nope/view").status_code == 404
    assert client.get("/documents/old/view").status_code == 404
