import io
import json
import mimetypes
import re
import shutil
import tempfile
import time
import uuid
from collections.abc import AsyncIterator
from contextlib import asynccontextmanager
from pathlib import Path
from typing import Literal

from fastapi import BackgroundTasks, Depends, FastAPI, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse, Response, StreamingResponse
from neo4j.exceptions import ConstraintError
import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from pydantic import BaseModel

from app import auth, model_settings, reindex, store, table
from app.config import settings
from app.graph import MAX_HISTORY_TURNS, to_citations, workflow
from app.ingest.convert import page_count, render_page
from app.ingest.pipeline import create_file, ingest, set_status
from app.models import chat_model, embedding_model
from app.table import (
    Column, ProposedColumns, Stopped, SuggestedTables, document_context, propose_columns, suggest_tables,
)
from app.table import run as run_table

READ_CHUNK = 1024 * 1024

CORS_ORIGINS = ["http://localhost:5173", "http://127.0.0.1:5173"]

@asynccontextmanager
async def lifespan(_: FastAPI) -> AsyncIterator[None]:
    # Background tasks live in-process, so a restart kills them; without this they read "running" forever
    try:
        store.fail_interrupted("interrupted by a server restart")
        auth.ensure_admin()
    except Exception as exc:
        # Neo4j down at boot: start anyway, the first request reports it
        print(f"could not finish startup: {exc}")
    yield


app = FastAPI(title="Docs Analyzer", lifespan=lifespan, dependencies=[Depends(auth.require_user)])
app.add_middleware(
    CORSMiddleware, allow_origins=CORS_ORIGINS, allow_credentials=True, allow_methods=["*"], allow_headers=["*"],
)


@app.exception_handler(Exception)
async def unhandled_exception_handler(request: Request, exc: Exception) -> JSONResponse:
    # A bare-Exception handler runs in ServerErrorMiddleware, outside CORSMiddleware, so it must add the
    # CORS headers itself or the browser sees an opaque CORS failure instead of the 500.
    # ponytail: exception text goes to the client as-is; sanitize and log before exposing beyond localhost
    origin = request.headers.get("origin")
    headers = (
        {"Access-Control-Allow-Origin": origin, "Access-Control-Allow-Credentials": "true"}
        if origin in CORS_ORIGINS else {}
    )
    return JSONResponse(status_code=500, content={"detail": str(exc)}, headers=headers)


class Login(BaseModel):
    username: str
    password: str


class User(BaseModel):
    id: str
    username: str
    role: str


class UserRow(User):
    created_at: int | None = None


class NewUser(BaseModel):
    username: str
    password: str
    role: Literal["admin", "user"] = "user"


class UserUpdate(BaseModel):
    password: str | None = None
    role: Literal["admin", "user"] | None = None


class PasswordChange(BaseModel):
    current: str
    new: str


Provider = Literal["ollama", "openai"]


class ModelSettings(BaseModel):
    """Keys go out masked; a key sent back as None or masked keeps the stored one."""
    llm_provider: Provider
    llm_base_url: str
    llm_api_key: str | None = None
    llm_model: str
    llm_model_choices: str = ""
    llm_reasoning_effort: str | None = None
    table_llm_base_url: str = ""
    table_llm_api_key: str | None = None
    table_llm_model: str = ""


class ModelSettingsView(ModelSettings):
    saved: bool


class EmbeddingSettings(BaseModel):
    embed_provider: Provider
    embed_base_url: str
    embed_api_key: str | None = None
    embed_model: str
    embed_dim: int
    reindex: bool = False


class Reindex(BaseModel):
    status: Literal["idle", "running", "done", "error"]
    done: int
    total: int
    error: str | None = None


class EmbeddingSettingsView(BaseModel):
    embed_provider: Provider
    embed_base_url: str
    embed_api_key: str
    embed_model: str
    embed_dim: int
    saved: bool
    chunks: int
    stored_model: str | None
    stored_dim: int | None
    reindex: Reindex


class ModelTest(BaseModel):
    slot: Literal["chat", "table", "embedding"]
    provider: Provider = "openai"
    base_url: str
    api_key: str | None = None
    model: str
    reasoning_effort: str | None = None


class ModelTestResult(BaseModel):
    ok: bool
    latency_ms: int
    reply: str | None = None
    dim: int | None = None
    error: str | None = None


class Limits(BaseModel):
    max_upload_mb: int
    max_pages: int


class Models(BaseModel):
    models: list[str]
    default: str


class Upload(BaseModel):
    doc_id: str


class Document(BaseModel):
    id: str
    name: str
    status: str
    progress: float
    created_at: int | None = None
    updated_at: int | None = None
    mode: str | None = None
    tokens: int | None = None
    chunks: int | None = None
    error: str | None = None
    step_done: int | None = None
    step_total: int | None = None
    llm_model: str | None = None
    language: str | None = None
    has_original: bool = False
    # Only on GET /documents/{id}: counting pages for the whole list would open every PDF
    page_count: int | None = None


class ProposeRequest(BaseModel):
    request: str


class TableRequest(BaseModel):
    request: str
    columns: list[Column]
    # Skip the scope pass and read every chunk
    read_all: bool = False


class Region(BaseModel):
    page: int
    l: float
    t: float
    r: float
    b: float
    page_width: float
    page_height: float


class Citation(BaseModel):
    n: int
    chunk_id: str
    pages: list[int]
    headings: list[str]
    source: str
    text: str
    regions: list[Region] = []


class Section(BaseModel):
    id: str
    title: str
    level: int
    order: int
    summary: str | None = None
    start_page: int | None = None
    end_page: int | None = None
    parent_id: str | None = None
    chunks: int


class Chunk(BaseModel):
    id: str
    idx: int
    text: str
    pages: list[int]
    headings: list[str]


class Row(BaseModel):
    id: str
    values: dict[str, str]
    citations: list[Citation]


class TableExtraction(BaseModel):
    id: str
    request: str
    columns: list[Column]
    rows: list[Row]
    status: str
    progress: float
    error: str | None = None
    chunks_done: int | None = None
    chunks_total: int | None = None
    read_parts: list[str] | None = None
    document_chunks: int | None = None
    read_all: bool = False
    created_at: int
    updated_at: int


class Turn(BaseModel):
    id: str
    question: str
    answer: str
    citations: list[Citation]
    mode: str
    complete: bool
    created_at: int


class GraphNode(BaseModel):
    key: str
    name: str
    type: str | None = None
    mentions: int


class GraphEdge(BaseModel):
    source: str
    target: str
    type: str


class GraphData(BaseModel):
    nodes: list[GraphNode]
    edges: list[GraphEdge]
    total: int


class EntityDetail(BaseModel):
    entity: GraphNode
    nodes: list[GraphNode]
    edges: list[GraphEdge]
    citations: list[Citation]


@app.post("/auth/login", response_model=User)
def login(body: Login, response: Response) -> dict:
    user = auth.authenticate(body.username, body.password)
    if user is None:
        raise HTTPException(status_code=401, detail="wrong username or password")
    response.set_cookie(
        auth.COOKIE, auth.start_session(user["id"]), max_age=auth.SESSION_DAYS * 86400,
        httponly=True, samesite="lax", secure=settings.cookie_secure,
    )
    return user


@app.post("/auth/logout", status_code=204)
def logout(request: Request, response: Response) -> None:
    auth.end_session(request.cookies[auth.COOKIE])
    response.delete_cookie(auth.COOKIE, httponly=True, samesite="lax", secure=settings.cookie_secure)


@app.get("/auth/me", response_model=User)
def me(user: dict = Depends(auth.require_user)) -> dict:
    return user


MIN_PASSWORD = 8
USERNAME = re.compile(r"[a-z0-9._-]{1,64}")


def checked_password(password: str) -> str:
    if len(password) < MIN_PASSWORD:
        raise HTTPException(status_code=400, detail=f"password must be at least {MIN_PASSWORD} characters")
    return password


@app.post("/auth/password", status_code=204)
def change_password(body: PasswordChange, request: Request, user: dict = Depends(auth.require_user)) -> None:
    if auth.authenticate(user["username"], body.current) is None:
        raise HTTPException(status_code=403, detail="current password is wrong")
    auth.set_password(user["id"], checked_password(body.new), keep_token=request.cookies[auth.COOKIE])


@app.get("/users", response_model=list[UserRow], dependencies=[Depends(auth.require_admin)])
def list_users() -> list[dict]:
    return auth.list_users()


@app.post("/users", response_model=User, dependencies=[Depends(auth.require_admin)])
def create_user(body: NewUser) -> dict:
    username = auth.normalize(body.username)
    if not USERNAME.fullmatch(username):
        raise HTTPException(status_code=400, detail="username: up to 64 of a-z, 0-9, dot, dash, underscore")
    try:
        return auth.create_user(username, checked_password(body.password), body.role)
    except ConstraintError:
        raise HTTPException(status_code=409, detail="username already taken")


def other_user(user_id: str, admin: dict) -> dict:
    # An admin never deletes or demotes themselves, so at least one admin always remains
    if user_id == admin["id"]:
        raise HTTPException(status_code=409, detail="not on your own account")
    user = auth.get_user(user_id)
    if user is None:
        raise HTTPException(status_code=404, detail="user not found")
    return user


@app.patch("/users/{user_id}", status_code=204)
def update_user(user_id: str, body: UserUpdate, admin: dict = Depends(auth.require_admin)) -> None:
    other_user(user_id, admin)
    if body.password is not None:
        auth.set_password(user_id, checked_password(body.password))
    if body.role is not None:
        auth.set_role(user_id, body.role)


@app.delete("/users/{user_id}", status_code=204)
def delete_user(user_id: str, admin: dict = Depends(auth.require_admin)) -> None:
    other_user(user_id, admin)
    auth.delete_user(user_id)


def get_file(doc_id: str) -> dict:
    file = store.get_file(doc_id)
    if file is None:
        raise HTTPException(status_code=404, detail="document not found")
    return file


def chat_history(doc_id: str) -> list[dict]:
    return store.recent_turns(doc_id, MAX_HISTORY_TURNS)


def original_path(doc_id: str) -> Path | None:
    # "{id}*", not "{id}.*": an Original uploaded without an extension is stored as the bare id
    return next((settings.data_dir / "originals").glob(f"{doc_id}*"), None)


def save_upload(file: UploadFile) -> str:
    """Stream the upload to a temp file, stopping as soon as it passes the size cap."""
    limit = settings.max_upload_mb * 1024 * 1024
    # Docling picks the format from the suffix
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=Path(file.filename or "").suffix)
    try:
        written = 0
        with tmp:
            while chunk := file.file.read(READ_CHUNK):
                written += len(chunk)
                if written > limit:
                    raise HTTPException(status_code=413, detail=f"file is larger than {settings.max_upload_mb} MB")
                tmp.write(chunk)
    except BaseException:
        Path(tmp.name).unlink(missing_ok=True)
        raise
    return tmp.name


@app.post("/documents", response_model=Upload)
def upload_document(file: UploadFile, background: BackgroundTasks) -> dict:
    no_reindex_running()
    path = save_upload(file)
    # Page cap is PDF-only: page_count reads PDFs only
    if path.lower().endswith(".pdf"):
        try:
            pages = page_count(path)
        except Exception:
            pages = 0
        if not pages:
            Path(path).unlink(missing_ok=True)
            raise HTTPException(status_code=400, detail="not a readable PDF")
        if pages > settings.max_pages:
            Path(path).unlink(missing_ok=True)
            raise HTTPException(status_code=413, detail=f"PDF has {pages} pages, over the {settings.max_pages} page limit")

    doc_id = create_file(file.filename or Path(path).name)
    originals = settings.data_dir / "originals"
    originals.mkdir(parents=True, exist_ok=True)
    original = originals / f"{doc_id}{Path(path).suffix.lower()}"
    shutil.move(path, original)
    background.add_task(ingest, doc_id, str(original))
    return {"doc_id": doc_id}


def model_settings_view() -> dict:
    values = model_settings.current()
    for field in model_settings.SECRETS:
        values[field] = model_settings.masked(values[field])
    return {**values, "saved": bool(model_settings.saved("models"))}


def checked_url(url: str, what: str) -> str:
    url = url.strip().rstrip("/")
    if not url.startswith(("http://", "https://")):
        raise HTTPException(status_code=400, detail=f"{what}: a URL starting with http:// or https://")
    return url


@app.get("/settings/models", response_model=ModelSettingsView, dependencies=[Depends(auth.require_admin)])
def get_model_settings() -> dict:
    return model_settings_view()


@app.put("/settings/models", response_model=ModelSettingsView, dependencies=[Depends(auth.require_admin)])
def put_model_settings(body: ModelSettings) -> dict:
    stored = model_settings.current()
    values = {k: v.strip() if isinstance(v, str) else v for k, v in body.model_dump().items()}
    for field in ("llm_api_key", "table_llm_api_key"):
        # A masked key echoed back is the stored key, never a new one
        if values[field] is None or values[field].startswith(model_settings.MASK):
            values[field] = stored[field]
    values["llm_base_url"] = checked_url(values["llm_base_url"], "chat model")
    if not values["llm_model"]:
        raise HTTPException(status_code=400, detail="chat model: a model name is required")
    if values["llm_provider"] == "openai" and not values["llm_api_key"]:
        raise HTTPException(status_code=400, detail="chat model: an OpenAI-compatible server needs an API key")
    if values["table_llm_model"]:
        values["table_llm_base_url"] = checked_url(values["table_llm_base_url"], "table model")
        if not values["table_llm_api_key"]:
            raise HTTPException(status_code=400, detail="table model: an API key is required")
    model_settings.save("models", values)
    return model_settings_view()


@app.delete("/settings/models", response_model=ModelSettingsView, dependencies=[Depends(auth.require_admin)])
def reset_model_settings() -> dict:
    model_settings.reset("models")
    return model_settings_view()


TEST_TIMEOUT_S = 30


@app.post("/settings/models/test", response_model=ModelTestResult, dependencies=[Depends(auth.require_admin)])
def test_model(body: ModelTest) -> dict:
    key = body.api_key
    if key is None or key.startswith(model_settings.MASK):
        key = model_settings.current()[{"chat": "llm_api_key", "table": "table_llm_api_key", "embedding": "embed_api_key"}[body.slot]]
    provider = "openai" if body.slot == "table" else body.provider
    started = time.monotonic()
    try:
        if body.slot == "embedding":
            vector = embedding_model(provider, checked_url(body.base_url, "URL"), key, body.model.strip()).embed_query("OK")
            return {"ok": True, "latency_ms": int((time.monotonic() - started) * 1000), "dim": len(vector)}
        client = chat_model(
            provider, checked_url(body.base_url, "URL"), key, body.model.strip(),
            max_tokens=256, reasoning_effort=body.reasoning_effort or None, timeout=TEST_TIMEOUT_S,
        )
        reply = client.invoke("Reply with the single word OK.").text
        return {"ok": True, "latency_ms": int((time.monotonic() - started) * 1000), "reply": reply.strip()[:200]}
    except HTTPException:
        raise
    except Exception as exc:
        return {"ok": False, "latency_ms": int((time.monotonic() - started) * 1000), "error": str(exc)[:500]}


def embedding_view() -> dict:
    values = model_settings.current()
    stored = reindex.marker() or {}
    return {
        **{field: values[field] for field in model_settings.GROUPS["embedding_endpoint"]},
        "embed_api_key": model_settings.masked(values["embed_api_key"]),
        "saved": bool(model_settings.saved("embedding_endpoint")),
        "chunks": reindex.embedded_chunks(),
        "stored_model": stored.get("model"),
        "stored_dim": stored.get("dim"),
        "reindex": reindex.STATE,
    }


@app.get("/settings/embedding", response_model=EmbeddingSettingsView, dependencies=[Depends(auth.require_admin)])
def get_embedding_settings() -> dict:
    return embedding_view()


@app.put("/settings/embedding", response_model=EmbeddingSettingsView, dependencies=[Depends(auth.require_admin)])
def put_embedding_settings(body: EmbeddingSettings, background: BackgroundTasks) -> dict:
    """Same model, or no vectors stored yet: saves at once. Another model over stored vectors needs a re-index."""
    if reindex.running():
        raise HTTPException(status_code=409, detail="a re-index is already running")
    values = {k: v.strip() if isinstance(v, str) else v for k, v in body.model_dump(exclude={"reindex"}).items()}
    if values["embed_api_key"] is None or values["embed_api_key"].startswith(model_settings.MASK):
        values["embed_api_key"] = model_settings.current()["embed_api_key"]
    values["embed_base_url"] = checked_url(values["embed_base_url"], "embedding model")
    if not values["embed_model"] or values["embed_dim"] <= 0:
        raise HTTPException(status_code=400, detail="embedding model: a model name and a positive dimension are required")
    if values["embed_provider"] == "openai" and not values["embed_api_key"]:
        raise HTTPException(status_code=400, detail="embedding model: an OpenAI-compatible server needs an API key")

    stored = reindex.marker() or {}
    same_model = stored.get("model") == values["embed_model"] and stored.get("dim") == values["embed_dim"]
    if same_model or not reindex.embedded_chunks():
        model_settings.save("embedding_endpoint", values)
        reindex.set_marker(values)
        return embedding_view()

    if not body.reindex:
        raise HTTPException(status_code=409, detail="another embedding model needs a re-index of every stored chunk")
    if reindex.ingests_in_flight():
        raise HTTPException(status_code=409, detail="wait for the documents still being ingested, then re-index")
    try:
        dim = reindex.probe(values)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"embedding model unreachable: {exc}")
    if dim != values["embed_dim"]:
        raise HTTPException(status_code=400, detail=f"{values['embed_model']} returns {dim} dimensions, not {values['embed_dim']}")
    if not reindex.start(values):
        raise HTTPException(status_code=409, detail="a re-index is already running")
    background.add_task(reindex.run, values)
    return embedding_view()


def no_reindex_running() -> None:
    # Chunks written mid re-index would keep the old model's vectors after the switch
    if reindex.running():
        raise HTTPException(status_code=409, detail="the embedding model is being switched, try again when it is done")


@app.get("/limits", response_model=Limits)
def limits() -> dict:
    return {"max_upload_mb": settings.max_upload_mb, "max_pages": settings.max_pages}


@app.get("/models", response_model=Models)
def models() -> dict:
    return {"models": settings.available_llm_models(), "default": settings.llm_model}


@app.get("/documents", response_model=list[Document])
def list_documents() -> list[dict]:
    return store.list_files()


@app.get("/documents/{doc_id}", response_model=Document)
def document_status(doc_id: str) -> dict:
    file = get_file(doc_id)
    path = original_path(doc_id)
    is_pdf = path is not None and mimetypes.guess_type(path)[0] == "application/pdf"
    return {**file, "has_original": path is not None, "page_count": page_count(path) if is_pdf else None}


@app.get("/documents/{doc_id}/original")
def document_original(doc_id: str) -> FileResponse:
    file = get_file(doc_id)
    path = original_path(doc_id)
    if path is None:
        raise HTTPException(status_code=404, detail="original not kept for this document")
    media_type = mimetypes.guess_type(path)[0] or "application/octet-stream"
    # Only PDFs open inline: an uploaded HTML must never run script on the API's origin
    disposition = "inline" if media_type == "application/pdf" else "attachment"
    return FileResponse(path, media_type=media_type, filename=file["name"], content_disposition_type=disposition)


@app.get("/documents/{doc_id}/original/pages/{page_no}")
def document_original_page(doc_id: str, page_no: int) -> Response:
    get_file(doc_id)
    path = original_path(doc_id)
    if path is None:
        raise HTTPException(status_code=404, detail="original not kept for this document")
    if mimetypes.guess_type(path)[0] != "application/pdf":
        raise HTTPException(status_code=404, detail="original is not a PDF")
    png = render_page(path, page_no)
    if png is None:
        raise HTTPException(status_code=404, detail="page not found")
    # An Original never changes and doc_id is fresh per upload, so cache forever
    return Response(png, media_type="image/png", headers={"Cache-Control": "public, max-age=31536000, immutable"})


@app.delete("/documents/{doc_id}", status_code=204)
def delete_document(doc_id: str) -> None:
    get_file(doc_id)
    # Safe mid-ingest: every ingest write is anchored on MATCH (f:File {id}), so nothing is orphaned.
    # Original first: on Windows unlink fails while ingest holds it open, and failing before the graph
    # delete leaves a Document the user can delete again instead of an orphaned file.
    if path := original_path(doc_id):
        path.unlink(missing_ok=True)
    store.delete_file(doc_id)


@app.post("/documents/{doc_id}/retry", response_model=Document)
def retry_document(doc_id: str, background: BackgroundTasks) -> dict:
    """Ingest again from the kept Original. Allowed on failed and done Documents, never on one still running."""
    file = get_file(doc_id)
    no_reindex_running()
    if file["status"] not in ("error", "done"):
        raise HTTPException(status_code=409, detail=f"the document is still being ingested (status: {file['status']})")
    path = original_path(doc_id)
    if path is None:
        raise HTTPException(status_code=409, detail="original not kept for this document, upload it again")
    store.reset_ingest(doc_id)
    background.add_task(ingest, doc_id, str(path))
    return {**get_file(doc_id), "has_original": True}


@app.get("/documents/{doc_id}/sections", response_model=list[Section])
def document_sections(doc_id: str) -> list[dict]:
    get_file(doc_id)
    return store.list_sections(doc_id)


@app.get("/documents/{doc_id}/sections/{section_id}/chunks", response_model=list[Chunk])
def document_section_chunks(doc_id: str, section_id: str) -> list[dict]:
    get_file(doc_id)
    return store.section_chunks(doc_id, section_id)


def extract_table(
    doc_id: str, table_id: str, request: str, columns: list[dict], model: str | None, read_all: bool = False,
) -> None:
    def progress(done: int, total: int) -> None:
        store.set_table_extraction(table_id, progress=round(done / total, 3), chunks_done=done, chunks_total=total)

    def cancelled() -> bool:
        extraction = store.get_table_extraction(doc_id, table_id)
        # Deleted counts as cancelled: nothing left to write the rows to
        return extraction is None or extraction["status"] == "cancelled"

    try:
        # Tables are read whole, not chunk by chunk (docs/adr/0005)
        tables = {t["id"]: t for t in store.list_doc_tables(doc_id)}
        document = [
            {**chunk, "tables": [tables[i] for i in chunk.get("table_ids") or [] if i in tables]}
            for chunk in store.list_chunks(doc_id)
        ]
        # docs/adr/0004
        chunks = document if read_all else table.scope(request, document, model)
        store.set_table_extraction(
            table_id, status="running", chunks_done=0, chunks_total=len(chunks),
            read_parts=table.part_titles(chunks), document_chunks=len(document),
        )
        rows = run_table(chunks, columns, model, on_progress=progress, should_stop=cancelled, request=request)
        store.set_table_extraction(
            # Row ids are positions at finish time; deleting a row never renumbers the rest
            table_id, status="done", progress=1.0, rows=[{**row, "id": str(i)} for i, row in enumerate(rows)],
            when_running=True,
        )
    except Stopped:
        store.set_table_extraction(table_id, status="cancelled")
    except Exception as exc:
        store.set_table_extraction(table_id, status="error", error=str(exc))


IN_FLIGHT = ("queued", "running")


def checked_columns(columns: list[Column]) -> list[dict]:
    # Names key each row's values: a blank or repeated name would silently drop a column's cells
    named = [{"name": c.name.strip(), "description": c.description.strip()} for c in columns]
    names = [c["name"] for c in named]
    if not names or not all(names):
        raise HTTPException(status_code=400, detail="every column needs a name")
    if len(set(names)) != len(names):
        raise HTTPException(status_code=400, detail="column names must be unique")
    return named


def ready_document(doc_id: str) -> dict:
    file = get_file(doc_id)
    if file["status"] != "done":
        raise HTTPException(status_code=409, detail=f"document is not ready (status: {file['status']})")
    return file


def table_context(doc_id: str) -> str:
    return document_context(store.list_sections(doc_id), store.list_chunks(doc_id, limit=table.CONTEXT_CHUNKS))


@app.post("/documents/{doc_id}/tables/propose", response_model=ProposedColumns)
def propose_table_columns(doc_id: str, body: ProposeRequest) -> dict:
    file = ready_document(doc_id)
    return {"columns": propose_columns(body.request, file.get("llm_model"), context=table_context(doc_id))}


@app.post("/documents/{doc_id}/tables/suggest", response_model=SuggestedTables)
def suggest_document_tables(doc_id: str) -> dict:
    # POST, not GET: an LLM call each time, so it runs only on request
    file = ready_document(doc_id)
    return {"tables": suggest_tables(table_context(doc_id), file.get("llm_model"))}


@app.post("/documents/{doc_id}/tables", response_model=TableExtraction)
def create_table_extraction(doc_id: str, body: TableRequest, background: BackgroundTasks) -> dict:
    file = ready_document(doc_id)
    columns = checked_columns(body.columns)
    # One run per document: each holds a worker thread for minutes
    if any(e["status"] in IN_FLIGHT for e in store.list_table_extractions(doc_id)):
        raise HTTPException(status_code=409, detail="a table extraction is already running for this document")
    extraction = store.create_table_extraction(doc_id, body.request, columns, read_all=body.read_all)
    background.add_task(extract_table, doc_id, extraction["id"], body.request, columns, file.get("llm_model"), body.read_all)
    return extraction


@app.get("/documents/{doc_id}/tables", response_model=list[TableExtraction])
def list_table_extractions(doc_id: str) -> list[dict]:
    get_file(doc_id)
    return store.list_table_extractions(doc_id)


@app.get("/documents/{doc_id}/tables/{table_id}", response_model=TableExtraction)
def document_table_extraction(doc_id: str, table_id: str) -> dict:
    get_file(doc_id)
    extraction = store.get_table_extraction(doc_id, table_id)
    if extraction is None:
        raise HTTPException(status_code=404, detail="table extraction not found")
    return extraction


@app.post("/documents/{doc_id}/tables/{table_id}/cancel", response_model=TableExtraction)
def cancel_table_extraction(doc_id: str, table_id: str) -> dict:
    extraction = document_table_extraction(doc_id, table_id)
    if extraction["status"] not in IN_FLIGHT:
        raise HTTPException(status_code=409, detail=f"table extraction is already {extraction['status']}")
    store.set_table_extraction(table_id, status="cancelled")
    return {**extraction, "status": "cancelled"}


@app.delete("/documents/{doc_id}/tables/{table_id}", status_code=204)
def delete_document_table_extraction(doc_id: str, table_id: str) -> None:
    # Fine mid-run: extract_table counts a missing run as cancelled
    document_table_extraction(doc_id, table_id)
    store.delete_table_extraction(doc_id, table_id)


@app.delete("/documents/{doc_id}/tables/{table_id}/rows/{row_id}", status_code=204)
def delete_table_row(doc_id: str, table_id: str, row_id: str) -> None:
    extraction = document_table_extraction(doc_id, table_id)
    if extraction["status"] in IN_FLIGHT:
        raise HTTPException(status_code=409, detail="the table extraction is still running")
    rows = [row for row in extraction["rows"] if row["id"] != row_id]
    if len(rows) == len(extraction["rows"]):
        raise HTTPException(status_code=404, detail="row not found")
    # ponytail: read-modify-write, two simultaneous deletes can lose one; edit in Cypher if that bites
    store.set_table_extraction(table_id, rows=rows)


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def cells(values: list) -> list:
    # Control characters make openpyxl refuse the whole workbook
    return [ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v for v in values]


@app.get("/documents/{doc_id}/tables/{table_id}/export.xlsx")
def export_table_extraction(doc_id: str, table_id: str) -> Response:
    extraction = document_table_extraction(doc_id, table_id)
    if extraction["status"] != "done":
        raise HTTPException(status_code=409, detail=f"table extraction is {extraction['status']}, not done")
    names = [column["name"] for column in extraction["columns"]]
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(cells([*names, "Pages"]))
    for row in extraction["rows"]:
        pages = sorted({page for citation in row["citations"] for page in citation["pages"]})
        sheet.append(cells([*(row["values"].get(name, "") for name in names), ", ".join(map(str, pages)) or None]))
    # Values come from the uploaded document: a cell reading "=HYPERLINK(...)" must stay text, not a formula
    for sheet_row in sheet.iter_rows():
        for cell in sheet_row:
            if cell.data_type == "f":
                cell.data_type = "s"
    buffer = io.BytesIO()
    workbook.save(buffer)
    # Headers must be latin-1, and the request can be in any script
    filename = " ".join(re.sub(r"[^A-Za-z0-9 _-]", " ", extraction["request"]).split())[:80] or "table"
    return Response(
        buffer.getvalue(), media_type=XLSX,
        headers={"Content-Disposition": f'attachment; filename="{filename}.xlsx"'},
    )


# ponytail: only the most mentioned entities, the rest via neighbour expansion; paginate or cluster if needed
GRAPH_NODES = 120
ENTITY_NEIGHBOURS = 40
ENTITY_EVIDENCE = 20


@app.get("/documents/{doc_id}/graph", response_model=GraphData)
def document_graph(doc_id: str, limit: int = GRAPH_NODES) -> dict:
    get_file(doc_id)
    return store.graph_top(doc_id, min(limit, 500))


@app.get("/documents/{doc_id}/graph/entity", response_model=EntityDetail)
def document_graph_entity(doc_id: str, key: str) -> dict:
    # key is a query parameter, not a path segment: entity names contain "/"
    get_file(doc_id)
    entity = store.entity(doc_id, key)
    if entity is None:
        raise HTTPException(status_code=404, detail="entity not found")
    neighbours = store.entity_neighbours(doc_id, key, ENTITY_NEIGHBOURS)
    nodes = list({n["key"]: {k: n[k] for k in ("key", "name", "type", "mentions")} for n in neighbours}.values())
    edges = [
        {"source": key, "target": n["key"], "type": n["rel"]} if n["outgoing"]
        else {"source": n["key"], "target": key, "type": n["rel"]}
        for n in neighbours
    ]
    chunks = store.entity_chunks(doc_id, key, ENTITY_EVIDENCE)
    return {
        "entity": entity, "nodes": nodes, "edges": edges,
        "citations": to_citations([{**chunk, "source": "kg"} for chunk in chunks]),
    }


@app.get("/documents/{doc_id}/turns", response_model=list[Turn])
def document_turns(doc_id: str) -> list[dict]:
    get_file(doc_id)
    return store.list_turns(doc_id)


@app.delete("/documents/{doc_id}/turns", status_code=204)
def clear_document_turns(doc_id: str) -> None:
    get_file(doc_id)
    store.clear_turns(doc_id)


def sse(event: str, data: object) -> str:
    # JSON even for plain tokens: a raw "\n" would split the "data:" line and corrupt the frame
    return f"event: {event}\ndata: {json.dumps(data)}\n\n"


async def stream_answer(
    doc_id: str, question: str, mode: str, model: str | None = None, history: list[dict] | None = None,
    language: str | None = None,
) -> AsyncIterator[str]:
    # Must stay an async generator: StreamingResponse runs a sync one in a worker thread that a client
    # disconnect never interrupts, so the finally block below would not run and no Turn gets saved.
    turn_id = str(uuid.uuid4())
    answer_parts: list[str] = []
    citations: list[dict] = []
    complete = False
    error: str | None = None
    try:
        async for kind, payload in workflow.astream(
            {"question": question, "doc_id": doc_id, "mode": mode, "model": model, "history": history or [],
             "language": language},
            stream_mode=["updates", "messages"],
        ):
            if kind == "updates":
                node, output = next(iter(payload.items()))
                if node == "fuse":
                    # Contexts are final here, so citations go out before the first answer token
                    citations = to_citations(output["contexts"])
                    yield sse("citations", citations)
                elif node != "generate":
                    yield sse("status", {"node": node})
            elif kind == "messages":
                chunk, meta = payload
                # Other nodes' structured-output calls stream too; without this their raw JSON leaks into the answer
                if meta.get("langgraph_node") != "generate" or not chunk.text:
                    continue
                answer_parts.append(chunk.text)
                yield sse("token", chunk.text)
        complete = True
    except Exception as exc:
        error = str(exc)
    finally:
        # Also runs on client disconnect (CancelledError is not an Exception), saving the partial answer
        # with complete=False
        store.save_turn(
            doc_id, question=question, answer="".join(answer_parts),
            citations=citations, mode=mode, complete=complete, turn_id=turn_id,
        )
    # After the save: the client refetches history on either event. "failed", not "error": EventSource
    # reserves "error" for its own connection failures.
    yield sse("failed", {"detail": error}) if error is not None else sse("done", {"turn_id": turn_id})


@app.get("/chat/stream")
def chat_stream(doc_id: str, question: str, model: str | None = None) -> StreamingResponse:
    file = get_file(doc_id)
    if file["status"] != "done":
        raise HTTPException(status_code=409, detail=f"document is not ready (status: {file['status']})")
    if model is not None and model not in settings.available_llm_models():
        raise HTTPException(status_code=400, detail=f"unknown model: {model}")
    # The picked model sticks to the document
    if model is not None and model != file.get("llm_model"):
        set_status(doc_id, llm_model=model)
    return StreamingResponse(
        stream_answer(
            doc_id, question, file["mode"], model or file.get("llm_model"), chat_history(doc_id), file.get("language"),
        ),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )
